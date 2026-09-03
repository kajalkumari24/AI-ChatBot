from __future__ import annotations

import csv
import io
import re
from datetime import datetime
from typing import Any


def _clean_filename(value: str) -> str:
   

    if not value:
        value = "generated_file"

    value = str(value).strip()

    # Remove characters that are unsafe in filenames.
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", value)

    # Replace multiple spaces.
    value = re.sub(r"\s+", "_", value)

    # Limit filename length.
    value = value[:100]

    return value or "generated_file"


def _timestamp() -> str:
    

    return datetime.now().strftime("%Y%m%d_%H%M%S")


def _result(
    filename: str,
    mime_type: str,
    file_bytes: bytes,
) -> dict[str, Any]:
    
    return {
        "filename": filename,
        "mime_type": mime_type,
        "file_bytes": file_bytes,
    }


def _normalize_text(content: Any) -> str:
   
    if content is None:
        return ""

    if isinstance(content, str):
        return content

    return str(content)



def generate_txt(
    content: str,
    filename: str | None = None,
) -> dict[str, Any]:
    

    text = _normalize_text(content)

    if filename:
        filename = _clean_filename(filename)
        if not filename.lower().endswith(".txt"):
            filename += ".txt"
    else:
        filename = f"generated_{_timestamp()}.txt"

    file_bytes = text.encode("utf-8")

    return _result(
        filename=filename,
        mime_type="text/plain; charset=utf-8",
        file_bytes=file_bytes,
    )


def generate_csv(
    data: Any,
    filename: str | None = None,
) -> dict[str, Any]:
    
    output = io.StringIO(newline="")

    writer = csv.writer(output)

    
    if isinstance(data, dict):

        writer.writerow(list(data.keys()))
        writer.writerow(list(data.values()))

    
    elif isinstance(data, (list, tuple)):

        if not data:
            writer.writerow([])

        elif all(isinstance(item, dict) for item in data):

            fieldnames: list[str] = []

            for item in data:
                for key in item.keys():
                    key = str(key)

                    if key not in fieldnames:
                        fieldnames.append(key)

            writer.writerow(fieldnames)

            for item in data:
                writer.writerow(
                    [
                        item.get(field, "")
                        for field in fieldnames
                    ]
                )

        elif all(
            isinstance(item, (list, tuple))
            for item in data
        ):

            for row in data:
                writer.writerow(row)

        else:

            for item in data:
                writer.writerow([item])

    
    else:

        text = _normalize_text(data)

        for line in text.splitlines():
            writer.writerow([line])

    file_bytes = output.getvalue().encode("utf-8-sig")

    if filename:
        filename = _clean_filename(filename)

        if not filename.lower().endswith(".csv"):
            filename += ".csv"

    else:
        filename = f"generated_{_timestamp()}.csv"

    return _result(
        filename=filename,
        mime_type="text/csv",
        file_bytes=file_bytes,
    )


def generate_docx(
    content: str,
    filename: str | None = None,
) -> dict[str, Any]:
   

    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(
            "DOCX generation requires python-docx. "
            "Install it with: pip install python-docx"
        ) from exc

    document = Document()

    text = _normalize_text(content)

    # Preserve paragraphs from generated content.
    paragraphs = text.split("\n")

    for paragraph in paragraphs:
        document.add_paragraph(paragraph)

    output = io.BytesIO()

    document.save(output)

    output.seek(0)

    if filename:
        filename = _clean_filename(filename)

        if not filename.lower().endswith(".docx"):
            filename += ".docx"

    else:
        filename = f"generated_{_timestamp()}.docx"

    return _result(
        filename=filename,
        mime_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        file_bytes=output.getvalue(),
    )



def generate_xlsx(
    data: Any,
    filename: str | None = None,
) -> dict[str, Any]:
    

    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError(
            "XLSX generation requires openpyxl. "
            "Install it with: pip install openpyxl"
        ) from exc

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Generated"

    
    if isinstance(data, dict):

        worksheet.append(
            ["Field", "Value"]
        )

        for key, value in data.items():

            worksheet.append(
                [
                    str(key),
                    value,
                ]
            )

    
    elif isinstance(data, (list, tuple)):

        if not data:
            worksheet.append([])

        elif all(isinstance(item, dict) for item in data):

            fieldnames: list[str] = []

            for item in data:
                for key in item.keys():

                    key = str(key)

                    if key not in fieldnames:
                        fieldnames.append(key)

            worksheet.append(fieldnames)

            for item in data:

                worksheet.append(
                    [
                        item.get(field, "")
                        for field in fieldnames
                    ]
                )

        elif all(
            isinstance(item, (list, tuple))
            for item in data
        ):

            for row in data:
                worksheet.append(list(row))

        else:

            for item in data:
                worksheet.append([item])

   
    else:

        text = _normalize_text(data)

        for line in text.splitlines():
            worksheet.append([line])

    
    for column in worksheet.columns:

        max_length = 0

        column_letter = column[0].column_letter

        for cell in column:

            value = cell.value

            if value is None:
                continue

            max_length = max(
                max_length,
                len(str(value))
            )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(max_length + 2, 10),
            60
        )

    output = io.BytesIO()

    workbook.save(output)

    output.seek(0)

    if filename:
        filename = _clean_filename(filename)

        if not filename.lower().endswith(".xlsx"):
            filename += ".xlsx"

    else:
        filename = f"generated_{_timestamp()}.xlsx"

    return _result(
        filename=filename,
        mime_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        file_bytes=output.getvalue(),
    )



def generate_pdf(
    content: str,
    filename: str | None = None,
) -> dict[str, Any]:
    

    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import (
            SimpleDocTemplate,
            Paragraph,
            Spacer,
        )
        from reportlab.lib.units import mm
    except ImportError as exc:
        raise RuntimeError(
            "PDF generation requires reportlab. "
            "Install it with: pip install reportlab"
        ) from exc

    if filename:
        filename = _clean_filename(filename)

        if not filename.lower().endswith(".pdf"):
            filename += ".pdf"

    else:
        filename = f"generated_{_timestamp()}.pdf"

    output = io.BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=20 * mm,
        leftMargin=20 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
    )

    styles = getSampleStyleSheet()

    body_style = styles["BodyText"]

    text = _normalize_text(content)

    story = []

    
    for line in text.splitlines():

        if not line.strip():

            story.append(
                Spacer(
                    1,
                    8
                )
            )

            continue

        safe_line = (
            line
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

        story.append(
            Paragraph(
                safe_line,
                body_style
            )
        )

        story.append(
            Spacer(
                1,
                4
            )
        )

    if not story:

        story.append(
            Paragraph(
                "",
                body_style
            )
        )

    document.build(story)

    output.seek(0)

    return _result(
        filename=filename,
        mime_type="application/pdf",
        file_bytes=output.getvalue(),
    )



def generate_file(
    file_type: str,
    content: Any,
    filename: str | None = None,
) -> dict[str, Any]:
    

    normalized_type = (
        str(file_type)
        .strip()
        .lower()
        .lstrip(".")
    )

    if normalized_type in {
        "txt",
        "text",
    }:

        return generate_txt(
            content,
            filename,
        )

    if normalized_type == "csv":

        return generate_csv(
            content,
            filename,
        )

    if normalized_type in {
        "docx",
        "doc",
    }:

        return generate_docx(
            content,
            filename,
        )

    if normalized_type in {
        "xlsx",
        "xls",
        "excel",
    }:

        return generate_xlsx(
            content,
            filename,
        )

    if normalized_type == "pdf":

        return generate_pdf(
            content,
            filename,
        )

    raise ValueError(
        "Unsupported file type: "
        f"{file_type}. "
        "Supported types are: "
        "txt, csv, docx, xlsx, pdf."
    )



__all__ = [
    "generate_txt",
    "generate_csv",
    "generate_docx",
    "generate_xlsx",
    "generate_pdf",
    "generate_file",
]
